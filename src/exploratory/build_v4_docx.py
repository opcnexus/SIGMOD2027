import os
#!/usr/bin/env python3
"""Build the reviewed v4 draft: figure rendering + table renumbering + captions + Word reconstruction."""
import re, os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
FIG = os.path.join(ROOT, "reports", "figures")
os.makedirs(FIG, exist_ok=True)
MD = os.path.join(ROOT, "reports", "paper_sigmod2027_cn_v4.md")
DOCX = os.path.join(ROOT, "reports", "paper_sigmod2027_cn_v4.docx")

plt.rcParams["font.sans-serif"] = ["PingFang SC", "Hiragino Sans GB", "Arial Unicode MS"]
plt.rcParams["axes.unicode_minus"] = False
C1, C2, C3 = "#637052", "#E48312", "#4a6b78"

# ---------- Figure 1: root% profile across the six platforms (log axis) ----------
ds = ["Molweni", "IRC", "Reddit", "Hacker News", "Discord", "Slack*"]
vals = [11.40, 28.05, 1.07, 0.09, 82.93, 11.12]
colors = [C1, C1, C3, C3, C2, C2]
fig, ax = plt.subplots(figsize=(7.2, 3.2))
bars = ax.bar(range(6), vals, color=colors, width=0.62)
ax.set_yscale("log"); ax.set_ylim(0.05, 300)
ax.set_xticks(range(6)); ax.set_xticklabels(ds, fontsize=10)
ax.set_ylabel("root% / mention rate (log axis)", fontsize=10)
for b, v in zip(bars, vals):
    ax.text(b.get_x() + b.get_width()/2, v*1.15, f"{v}", ha="center", fontsize=9)
ax.axhline(1.0, ls="--", lw=0.8, color="#999")
ax.set_title("Representation profile across six platforms: root% spans three orders of magnitude (annotation tiers colour-coded)", fontsize=11)
plt.tight_layout(); plt.savefig(f"{FIG}/fig1_layering.png", dpi=200); plt.close()

# ---------- Figure 2: the three-stage pipeline ----------
fig, ax = plt.subplots(figsize=(7.2, 3.6)); ax.axis("off")
def box(x, y, w, h, text, c):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.06",
                                fc=c, ec="#333", lw=1.2))
    ax.text(x+w/2, y+h/2, text, ha="center", va="center", fontsize=10)
def arrow(x1, y1, x2, y2):
    ax.add_patch(FancyArrowPatch((x1, y1), (x2, y2), arrowstyle="-|>",
                                 mutation_scale=16, color="#333", lw=1.2))
box(0.02, 0.55, 0.30, 0.34, "Stage 1 Projection measurement\n6-dim profile of D\nO(n) - language-agnostic", C1)
box(0.36, 0.55, 0.30, 0.34, "Stage 2 Stratified construction\nStrategies A/B/C + gates\nG1 parse rate - G2 temporal - G3 multi-parent", C2)
box(0.70, 0.55, 0.28, 0.34, "Stage 3 Generalization diagnostics\nUnified window protocol\nCriteria J1/J2/J3", C3)
for x in (0.32, 0.66):
    arrow(x, 0.72, x+0.04, 0.72)
box(0.06, 0.10, 0.26, 0.26, "Forum tree-structured\nroot~0", "#e8f0e3")
box(0.38, 0.10, 0.26, 0.26, "Annotated conversational\nroot 11-28%", "#fdf1e3")
box(0.70, 0.10, 0.26, 0.26, "Real-world chat\nroot>=80%", "#e3edf0")
arrow(0.17, 0.55, 0.17, 0.38); arrow(0.50, 0.55, 0.50, 0.38); arrow(0.84, 0.55, 0.84, 0.38)
ax.set_xlim(0, 1); ax.set_ylim(0, 1)
ax.set_title("The three-stage RepSpace pipeline", fontsize=11)
plt.tight_layout(); plt.savefig(f"{FIG}/fig2_pipeline.png", dpi=200); plt.close()

# ---------- Figure 3: algorithm architecture ----------
fig, ax = plt.subplots(figsize=(7.2, 3.8)); ax.axis("off")
box(0.02, 0.70, 0.28, 0.22, "Unified input layer\n(author, ts, text, parents)\nchar-3gram->64d + 6-dim edge features", C1)
box(0.36, 0.74, 0.26, 0.16, "Encoding layer\nHasher / causal attention\nDistilBERT (layers 4-5 fine-tuned)", C2)
box(0.36, 0.46, 0.26, 0.16, "Scoring and decoding layer\nheuristics / PairMLP\nMHA-Net / DiGAT pointer", C3)
box(0.04, 0.30, 0.26, 0.20, "Heuristic probes\nprev1 . saprev . simmax\n(parameter-free)", "#e8f0e3")
box(0.36, 0.18, 0.26, 0.20, "Learned probes\nPairMLP . MHA-Net\n(5 seeds)", "#fdf1e3")
box(0.68, 0.30, 0.28, 0.20, "ECB event-chain reconstruction\nLLM eventization -> assembly\n-> backtracking completion (this paper)", "#e3edf0")
box(0.36, 0.02, 0.26, 0.10, "Unified evaluation over 12+3 metrics", "#f5e8e8")
for (x1,y1,x2,y2) in [(0.30,0.80,0.36,0.82),(0.49,0.74,0.49,0.62),
                      (0.30,0.40,0.36,0.40),(0.62,0.28,0.68,0.40),
                      (0.49,0.46,0.49,0.40),(0.62,0.12,0.49,0.02)]:
    arrow(x1,y1,x2,y2)
ax.set_xlim(0, 1); ax.set_ylim(0, 1)
ax.set_title("The calibrated probe family and the ECB architecture", fontsize=11)
plt.tight_layout(); plt.savefig(f"{FIG}/fig3_architecture.png", dpi=200); plt.close()

# ---------- table renumbering + synchronization of in-text references ----------
s = open(MD).read()
reps = [
    ("（见表 1）", "（见表 2）"),
    ("表 3 给出主诊断结果", "表 4 给出主诊断结果"),
    ("表 5 给出跨层级迁移矩阵", "表 6 给出跨层级迁移矩阵"),
    ("表 8 给出回溯重构准确率矩阵", "表 7 给出回溯重构准确率矩阵"),
    ("对照（表 10）：", "对照（表 9）："),
    ("判据适用性见表 6", "判据适用性见表 11"),
    ("**Table 1：六平台语料", "**Table 2：六平台语料"),
    ("**Table 2：与既有解缠语料", "**Table 3：与既有解缠语料"),
    ("**Table 3：主诊断表", "**Table 4：主诊断表"),
    ("**Table 4：按金标结构分解", "**Table 5：按金标结构分解"),
    ("**Table 5：跨层级迁移矩阵", "**Table 6：跨层级迁移矩阵"),
    ("**Table 8：M3 回溯重构准确率矩阵", "**Table 7：M3 回溯重构准确率矩阵"),
    ("**Table 9：OCR-APT 式剪枝", "**Table 8：OCR-APT 式剪枝"),
    ("**Table 10：三方链级对照", "**Table 9：三方链级对照"),
    ("**Table 7：路由原型结果", "**Table 10：路由原型结果"),
    ("**Table 6：向相邻任务的映射", "**Table 11：向相邻任务的映射"),
    ("六个平台在统一会话表示（作者、父边、时间戳、文本）下对齐。",
     "六个平台在统一会话表示（作者、父边、时间戳、文本）下对齐，其 root% 画像见图 1。"),
    ("工程上，数据管线形式化为四个幂等算子：",
     "工程上，数据管线形式化为四个幂等算子（算子定义见表 1，三阶段流程见图 2）："),
    ("判据不依赖特定算法族——本文以五个校准探针实例化，替换探针不影响判据的有效性。",
     "判据不依赖特定算法族——本文以五个校准探针实例化（探针族与 ECB 架构见图 3），替换探针不影响判据的有效性。"),
]
for a, b in reps:
    if a in s:
        s = s.replace(a, b)
# add a caption to the operator table
s = s.replace("| 算子 | 输入 | 输出 | 错误处理 |",
              "**Table 1：数据准备算子定义**\n\n| 算子 | 输入 | 输出 | 错误处理 |")
open(MD, "w").write(s)
print("md refs synced")

# ---------- table captions in document order ----------
CAPS = ["表 1：数据准备算子定义",
        "表 2：六平台语料库统计与质量门禁结果",
        "表 3：与既有解缠语料的互补性对比",
        "表 4：主诊断表（test 链接 F1，学习型为 5 seeds 均值±标准差）",
        "表 5：按金标结构分解的边类别 F1（各数据集最优探针，5 seeds 均值）",
        "表 6：跨层级迁移衰减矩阵（统一探针，test F1）",
        "表 7：M3 回溯重构准确率矩阵（test，%）",
        "表 8：OCR-APT 式剪枝对照（test 链级）",
        "表 9：三方链级对照（全量 test；ECB 均值）",
        "表 10：路由原型结果（test link-F1，探针族 prev1/saprev/simmax/PairMLP）",
        "表 11：向相邻任务的映射"]

# ---------- docx reconstruction (with figures) ----------
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin, sec.bottom_margin = Cm(2.54), Cm(2.54)

def set_font(run, cn="宋体", en="Times New Roman", size=10.5, bold=False):
    run.font.name = en; run.font.size = Pt(size); run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), cn)

def para(text, size=10.5, bold=False, cn="宋体", align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         indent=True, after=6):
    p = doc.add_paragraph(); p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if indent: p.paragraph_format.first_line_indent = Pt(size*2)
    set_font(p.add_run(text), cn, size=size, bold=bold)
    return p

def heading(text, lv):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if lv == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    set_font(p.add_run(text), "黑体", size={1: 15, 2: 12, 3: 11}.get(lv, 11), bold=True)

def caption(text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    set_font(p.add_run(text), "黑体", size=9.5, bold=True)

def add_table(rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = t.cell(ri, ci); c.text = ""
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(cell), "宋体", size=9, bold=(ri == 0))
            if ri == 0:
                sh = OxmlElement("w:shd"); sh.set(qn("w:fill"), "EFEFEF")
                c._tc.get_or_add_tcPr().append(sh)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

FIGS = {  # anchor heading -> (image, caption)
    "## 2 相关工作": ("fig1_layering.png", "图 1：六平台表征画像——root% 跨越三个数量级（对数轴；绿=标注会话型，蓝=论坛树型，橙=真实聊天型；Slack 为提及率）"),
    "### 4.3": ("fig2_pipeline.png", "图 2：RepSpace 三阶段流程——投影测量、分层构建与泛化诊断"),
    "## 5 实例化一": ("fig3_architecture.png", "图 3：校准探针族与 ECB 事件链重构的统一架构"),
}
cap_idx = 0
lines = open(MD).read().split("\n")
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    for anchor, (img, cap) in FIGS.items():
        if line.strip().startswith(anchor):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(os.path.join(FIG, img), width=Cm(15.5))
            caption(cap)
            break
    if not line.strip():
        i += 1; continue
    if line.startswith("# ") and i < 3:
        para(line[2:], size=16, bold=True, cn="黑体", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, after=4)
    elif line.startswith("## "):
        heading(line[3:], 1)
    elif line.startswith("### "):
        heading(line[4:], 2)
    elif line.startswith("|"):
        if cap_idx < len(CAPS):
            caption(CAPS[cap_idx]); cap_idx += 1
        tbl = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"-{2,}:?", c) for c in cells):
                tbl.append(cells)
            i += 1
        add_table(tbl)
        continue
    elif re.match(r"^\[\d+\]", line.strip()):
        p = para(line.strip(), size=9, indent=False)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.8)
    elif line.startswith("副标题") or line.startswith("（SIGMOD"):
        para(line, size=11, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    else:
        para(line)
    i += 1

doc.save(DOCX)
print("saved", DOCX, "| tables captioned:", cap_idx)
